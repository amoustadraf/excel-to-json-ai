import json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

# Hex validation
HEX6 = re.compile(r'^[0-9A-F]{6}$')
HEX8 = re.compile(r'^[0-9A-F]{8}$')

def sanitize_color(s):
    """Return 8-digit ARGB or None."""
    if not isinstance(s, str):
        return None
    c = s.lstrip('#').upper()
    if HEX6.match(c):
        return 'FF' + c
    if HEX8.match(c):
        return c
    return None

# Load data
with open("output.json", "r", encoding="utf-8") as f:
    data = json.load(f)

wb = Workbook()
wb.remove(wb.active)

for sheet_name, cells in data.items():
    ws = wb.create_sheet(title=sheet_name)
    for coord, info in cells.items():
        cell = ws[coord]
        cell.value = info.get("value")

        # colors
        font_color = sanitize_color(info.get("font_color"))
        fill_color = sanitize_color(info.get("fill_color"))

        # Font
        cell.font = Font(
            bold=info.get("bold", False),
            italic=info.get("italic", False),
            underline='single' if info.get("underline", False) else None,
            strikethrough=info.get("strikethrough", False), 
            color=font_color if font_color else None
        )

        # Fill (use fgColor for actual fill)
        if fill_color:
            cell.fill = PatternFill(
                patternType="solid",
                fgColor=fill_color
            )

        # Borders
        cell.border = Border(
            top=Side(style=info.get("border_top")),
            bottom=Side(style=info.get("border_bottom")),
            left=Side(style=info.get("border_left")),
            right=Side(style=info.get("border_right"))
        )

wb.save("sandbox.xlsx")